import requests
import openpyxl
from openpyxl import Workbook
from bs4 import BeautifulSoup
from datetime import datetime
from openpyxl.chart import LineChart, Reference

# Empty list to store the cleaned buy/sell prices
price_list = []

# Request the page and parse its HTML
web = requests.get("https://dolarhoy.com/cotizacion-dolar-blue")
html = BeautifulSoup(web.text, "html.parser")

# Buy and sell prices are inside <div class="value"> tags
prices = html.find_all("div", class_="value")

# Remove the leading "$" from each price with slicing and store the clean values
for price in prices:
    price_list.append(price.text[1:])

# Current date and time, formatted as day/month/year hour:minute
now = datetime.now().strftime("%d/%m/%y %H:%M")

# Try to open the existing tracker file. If it doesn't exist yet,
# create it and add the header row (this only runs once, the first time)
try:
    wb = openpyxl.load_workbook("Dolar Tracker.xlsx")
except FileNotFoundError:
    wb = Workbook()
    sheet = wb.active
    sheet.append(["FECHA", "COMPRA", "VENTA"])
    wb.save("Dolar Tracker.xlsx")

# Add a new row with the current timestamp and prices,
# whether the file was just created or already existed
sheet = wb.active
sheet.append([now, price_list[0], price_list[1]])

chart = LineChart()
data = Reference(sheet, min_col = 2, min_row = 1, max_row = sheet.max_row)
chart.add_data(data, titles_from_data = True)
sheet._charts = []
sheet.add_chart(chart, "E2")

wb.save("Dolar Tracker.xlsx")